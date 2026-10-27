from __future__ import annotations

import hashlib
import json
import tempfile
import unittest
from contextlib import ExitStack
from pathlib import Path
from unittest.mock import patch
from xml.etree import ElementTree as ET

from openpyxl import load_workbook

from term_extractor_app.ai_review.inline_xml import inner_xml
from term_extractor_app.ai_review.readers import XmlReader, XliffReader
from term_extractor_app.ai_review.xliff_reader import read_xliff_items
from term_extractor_app.ai_review.workspace_service import _build_deterministic_plan, _plan_from_agent
from term_extractor_app.ai_review import review_service as review


class InlineContentTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.root = Path(self.temp.name)

    def tearDown(self):
        self.temp.cleanup()

    def fixture(self, name, xml):
        path = self.root / name
        ET.ElementTree(ET.fromstring(xml)).write(path, encoding="utf-8", xml_declaration=True)
        return path

    def test_xliff_12_preserves_placeholders_attributes_nested_codes_and_tails(self):
        path = self.fixture("inline.xlf", '''<xliff xmlns="urn:oasis:names:tc:xliff:document:1.2" version="1.2">
        <file source-language="en" target-language="zh"><body><trans-unit id="1">
        <source>Hello <ph id="1" equiv-text="{player}"/>! <g id="2">Click <x id="3"/></g> now.</source>
        <target>你好！<g id="2">点击<x id="3"/></g>。</target>
        </trans-unit></body></file></xliff>''')
        item = read_xliff_items(path, "inline.xlf")[0]
        self.assertIn('<ph id="1" equiv-text="{player}"', item["source_text"])
        self.assertIn('<g id="2">Click <x id="3" /></g> now.', item["source_text"])
        self.assertNotIn("ns0:", item["source_text"])
        self.assertNotIn("{player}", item["target_text"])
        document = XliffReader().read(path, "inline.xlf")
        self.assertEqual(document.blocks[0].metadata["source_text"], item["source_text"])
        self.assertEqual(document.blocks[0].text, item["target_text"])

    def test_xliff_annotations_preserve_content_without_becoming_required_codes(self):
        element = ET.fromstring('<source>A<mrk mtype="seg">B<g id="2">C<mrk mtype="term">D</mrk>E</g>F</mrk>G<ph id="1"/>H</source>')
        original = ET.tostring(element)
        self.assertEqual(inner_xml(element, xliff=True), 'AB<g id="2">CDE</g>FG<ph id="1" />H')
        self.assertEqual(ET.tostring(element), original)

    def test_xliff_paired_native_codes_keep_encoded_original_markup(self):
        element = ET.fromstring('<source><bpt id="1">&lt;b&gt;</bpt>Bold<ept id="1">&lt;/b&gt;</ept> tail</source>')
        value = inner_xml(element, xliff=True)
        self.assertEqual(value, '<bpt id="1">&lt;b&gt;</bpt>Bold<ept id="1">&lt;/b&gt;</ept> tail')

    def test_xliff_20_referenced_original_data_reaches_request_payload(self):
        path = self.fixture("v2.xlf", '''<xliff xmlns="urn:oasis:names:tc:xliff:document:2.0" version="2.0" srcLang="en" trgLang="zh">
        <file id="f"><unit id="u"><originalData><data id="d1">{player}</data><data id="unused">large unrelated code</data></originalData>
        <segment id="s"><source>Hi <ph id="1" dataRef="d1"/></source><target>你好</target></segment></unit></file></xliff>''')
        item = read_xliff_items(path, "v2.xlf")[0]
        self.assertEqual(item["inline_code_data"], {"d1": "{player}"})
        document = XliffReader().read(path, "v2.xlf")
        session = {"id": "session", "source_language": "en", "target_languages": ["zh"]}
        plan = _build_deterministic_plan(session, {"a": document}, "")
        unit = plan.targets[0].units[0]
        payload = review._build_request_payload([{
            "id": "item", "source_text": unit.source_text, "target_text": unit.target_text,
            "info_json": json.dumps(unit.references, ensure_ascii=False),
        }], {})
        self.assertIn("{player}", payload["items"][0]["info"][0]["value"])
        self.assertNotIn("unused", payload["items"][0]["info"][0]["value"])
        agent_plan = _plan_from_agent("session", {"source_language": "en", "targets": [{
            "language": "zh", "mappings": [{"attachment_id": "a", "target_pointer": "segment:s"}]
        }]}, {"a": document})
        self.assertEqual(agent_plan.targets[0].units[0].references, unit.references)

    def test_xml_mixed_content_retains_markup_and_tail_once(self):
        path = self.fixture("mixed.xml", '<root><p>Click <b>here <i>now</i></b> to start.</p><p>Other text</p></root>')
        document = XmlReader().read(path, "mixed.xml")
        self.assertEqual([block.text for block in document.blocks], [
            'Click <b>here <i>now</i></b> to start.', 'Other text'])
        self.assertEqual([block.pointer for block in document.blocks], ['element:p[1]', 'element:p[2]'])

    def test_xml_leading_inline_child_and_tail_preserve_full_sentence(self):
        path = self.fixture("leading.xml", '<root><p><b>Start</b> here.</p><row><source>A</source><target>B</target></row></root>')
        self.assertEqual([block.text for block in XmlReader().read(path, "leading.xml").blocks], [
            '<b>Start</b> here.', 'A', 'B'])

    def test_inline_tag_reordering_is_preserved(self):
        source = inner_xml(ET.fromstring('<source><ph id="1"/> A <ph id="2"/></source>'), xliff=True)
        target = inner_xml(ET.fromstring('<target><ph id="2"/> B <ph id="1"/></target>'), xliff=True)
        self.assertLess(source.index('id="1"'), source.index('id="2"'))
        self.assertLess(target.index('id="2"'), target.index('id="1"'))

    def test_inline_codes_survive_session_storage_export_and_upload_cleanup(self):
        from term_extractor_app.ai_review import database as db, config, cache_service, session_store
        from term_extractor_app.ai_review import output_service, workflow_service, workspace_service
        path = self.fixture("pipeline.xlf", '''<xliff version="2.0" srcLang="en" trgLang="zh"><file id="f"><unit id="u">
        <originalData><data id="d1">{player}</data></originalData>
        <segment id="s"><source>Hello <ph id="1" dataRef="d1"/></source><target>你好</target></segment></unit></file></xliff>''')
        uploads = self.root / "uploads"
        uploads.mkdir()
        with ExitStack() as stack:
            for module, key, value in [
                (db, "DB_PATH", self.root / "review.sqlite3"),
                (config, "UPLOADS_DIR", uploads), (cache_service, "UPLOADS_DIR", uploads),
                (session_store, "UPLOADS_DIR", uploads), (output_service, "OUTPUTS_DIR", self.root),
            ]:
                stack.enter_context(patch.object(module, key, value))
            stack.enter_context(patch.object(review, "get_shared_ai_settings", return_value={"api_key": "dummy", "selected_model": "mock"}))
            stack.enter_context(patch.object(review, "_track_ai_review_start"))
            stack.enter_context(patch.object(review, "_add_log"))
            stack.enter_context(patch.object(review.threading, "Thread"))
            session = session_store.create_session(source_language="en", target_languages=["zh"])
            attachment = session_store.add_attachment(session["id"], "pipeline.xlf", path.read_bytes())
            workspace_service.inspect_workspace_sync(session["id"], attachment_ids=[attachment["id"]])
            task_id = workflow_service.start_session_review(session["id"])[0]["task_id"]
            task = review.get_review_task(task_id)
            items = review._get_batch_items(task["batch_id"])
            payload = review._build_request_payload(items, task["config"])
            self.assertIn('<ph id="1" dataRef="d1"', payload["items"][0]["source"])
            self.assertIn("{player}", payload["items"][0]["info"][0]["value"])
            review._save_review_results_bulk(task_id, [(items[0], {
                "id": items[0]["id"], "has_issue": False, "issue_type": "", "issue": "", "suggestion": "",
            })])
            exported = output_service.generate_review_excel(task_id)
            book = load_workbook(exported, read_only=True)
            try:
                row = next(book["审校结果"].iter_rows(min_row=2, values_only=True))
                self.assertIn('<ph id="1" dataRef="d1"', row[3])
            finally:
                book.close()
            review._update_task(task_id, status="completed", output_path=str(exported))
            session_store.update_session(session["id"], status="completed")
            summary = session_store.release_completed_session_cache(session["id"])
            self.assertEqual(summary["released_attachment_count"], 1)
            retained = review._get_batch_items(task["batch_id"])
            self.assertEqual(review._build_request_payload(retained, task["config"]), payload)


class EffectivePromptCacheTests(unittest.TestCase):
    def signature(self, memory):
        system = review._effective_system_prompt({"system_prompt": "Review.", "memory_snapshot": memory})
        return review._prompt_signature(system, "{text}", "en", "zh")

    def test_active_rule_and_example_changes_invalidate_identity(self):
        base = {"version": 1, "rules": [{"id": "a", "text": "Use context.", "active": True, "weight": 1}]}
        changed_rule = {"rules": [{**base["rules"][0], "text": "Check punctuation."}]}
        example = {"rules": [{**base["rules"][0], "few_shot": {"source_text": "Hello", "target_text": "你好"}}]}
        changed_example = {"rules": [{**example["rules"][0], "few_shot": {"source_text": "Hi", "target_text": "您好"}}]}
        self.assertNotEqual(self.signature(base), self.signature(changed_rule))
        self.assertNotEqual(self.signature(base), self.signature(example))
        self.assertNotEqual(self.signature(example), self.signature(changed_example))
        self.assertNotEqual(self.signature(base), self.signature({}))

    def test_inactive_rules_version_and_weights_do_not_invalidate_effective_prompt(self):
        rule = {"id": "a", "text": "Use context.", "active": True, "weight": 1}
        base = {"version": 1, "rules": [rule]}
        changed = {"version": 2, "rules": [{**rule, "weight": 2}, {"id": "b", "text": "Candidate", "active": False}]}
        self.assertEqual(self.signature(base), self.signature(changed))

    def test_old_cache_namespace_expires_without_deleting_history(self):
        old = json.dumps({"system_prompt": "Review.", "user_prompt": "{text}", "source_language": "en", "target_language": "zh", "term_base_hash": ""}, ensure_ascii=False)
        self.assertNotEqual(self.signature({}), hashlib.sha256(old.encode("utf-8")).hexdigest())


if __name__ == "__main__":
    unittest.main()
