from __future__ import annotations

import io
import tempfile
import time
import unittest
from pathlib import Path
from unittest.mock import patch

from fastapi import FastAPI
from fastapi.testclient import TestClient
from openpyxl import load_workbook

from term_extractor_app.ai_review import database as db, learning_service as learning, output_service as output, review_service as review
from term_extractor_app.ai_review.conversation_routes import router
from term_extractor_app.ai_review.session_store import create_session, delete_session


class ReviewLearningTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.root = Path(self.temp.name)
        self.patches = [patch.object(db, 'DB_PATH', self.root / 'review.db'),
                        patch.object(output, 'OUTPUTS_DIR', self.root),
                        patch.object(learning, 'start_learning'),
                        patch.object(review, '_add_log')]
        for p in self.patches:
            p.start()
        db.init_db()
        self.session = create_session(title='Learning QA')['id']
        self.other = create_session(title='Other QA')['id']
        self.config = {'mode':'normal', 'session_id':self.session, 'enable_ai_review':True,
                       'model':'mock', 'system_prompt':'review', 'user_prompt':'{text}',
                       'source_language':'英语', 'target_language':'简体中文',
                       'max_chars_per_request':3000}
        with db.get_connection() as conn:
            conn.execute("INSERT INTO file_batches(id, original_filename, stored_path, file_type, status, created_at, updated_at) VALUES('b','qa.xlsx','','xlsx','ready','now','now')")
            conn.execute("INSERT INTO review_tasks(id,batch_id,session_id,status,config_json,created_at,updated_at) VALUES('t','b',?,'completed',?,'now','now')", (self.session, db.dumps_json(self.config)))
            conn.execute("INSERT INTO review_session_tasks VALUES(?,'简体中文','t','now')", (self.session,))
            for i in range(3):
                conn.execute("INSERT INTO file_items(id,batch_id,source_file,source_text,target_text,item_order,created_at) VALUES(?,'b','qa.xlsx',?,?,?,'now')", (f'i{i}',f'source{i}',f'target{i}',i))
        self.items = [{'id':f'i{i}', 'source_text':f'source{i}', 'target_text':f'target{i}', 'cache_key':f'k{i}'} for i in range(3)]
        review._save_review_results_bulk('t', [(item, {'id':item['id'],'has_issue': True, 'issue':'误报', 'issue_type':'术语','suggestion':'suggestion'}) for item in self.items])
        self.results = review.get_review_results('t', None)
        app = FastAPI(); app.include_router(router)
        self.client = TestClient(app)

    def tearDown(self):
        self.client.close()
        for p in reversed(self.patches): p.stop()
        self.temp.cleanup()

    def feedback(self, index=0, reason='符合上下文'):
        return learning.submit_feedback(self.session, 't', [{'result_id':self.results[index]['id'], 'decision':'disagree', 'reason':reason}])

    def test_defaults_and_settings_api(self):
        url='/api/ai-review/conversations/learning/options'
        self.assertEqual(self.client.get(url).json(), {'cache_enabled':True,'learning_enabled':True})
        self.assertEqual(self.client.post(url, json={'cache_enabled':False,'learning_enabled':False}).status_code, 200)
        self.assertFalse(self.client.get(url).json()['cache_enabled'])

    def test_remote_lookup_covers_all_segments(self):
        from term_extractor_app.ai_review.memoq_service import MemoQClient
        client = MemoQClient({'username':'qa','password':'qa'})
        try:
            with patch.object(client, 'request', return_value={}) as call:
                client.lookup(['a','b'], 'eng', 'jpn', [str(i) for i in range(65)])
                sizes=[len(c.kwargs['json']['Segments']) for c in call.call_args_list]
                self.assertEqual(sizes, [32,32,1,32,32,1])
        finally:
            client.close()

    def test_remote_terms_only_include_current_package(self):
        pairs=[{'source':'source0','targets':['a']},{'source':'source1','targets':['b']}]
        payload=review._build_request_payload([self.items[0]], {'memoq_term_pairs':pairs})
        self.assertEqual(payload['term_pairs'],pairs[:1])

    def test_prompt_and_term_note_changes_invalidate_signature(self):
        first=review._prompt_signature('system','{text}','eng','jpn','revision1')
        self.assertNotEqual(first,review._prompt_signature('changed','{text}','eng','jpn','revision1'))
        self.assertNotEqual(first,review._prompt_signature('system','{text}','eng','jpn','revision2'))
        args=['source','target',[],[{'source':'source','targets':['target'],'entry_note':'a'}],'model',first,'',False]
        key=review._cache_key(*args)
        args[3][0]['entry_note']='b'
        self.assertNotEqual(key,review._cache_key(*args))

    def test_memory_injection_only_active_rules(self):
        self.assertEqual(learning.memory_prompt({}), '')
        text=learning.memory_prompt({'rules':[{'text':'active norm','active':True},{'text':'candidate norm','active':False}]})
        self.assertIn('active norm',text); self.assertNotIn('candidate norm',text)

    def test_few_shot_soft_length_limit_provenance_and_no_truncation(self):
        feedback = [{'source_text':'s' * 101, 'target_text':'译' * 101},
                    {'source_text':'other', 'target_text':'different'}]
        def apply(example):
            return learning.apply_memory_update([], {'triggered_ids':[], 'new_rules':[
                {'text':'norm', 'few_shot':example}]}, feedback)
        valid = {'feedback_index':0, 'source_text':'s' * 100, 'target_text':'译' * 100}
        self.assertEqual(len(apply(valid)[0]['few_shot']['target_text']), 100)
        long_example = {'feedback_index':0, 'source_text':'s' * 101, 'target_text':'译' * 101}
        saved = apply(long_example)[0]['few_shot']
        self.assertEqual(saved, {key:long_example[key] for key in ('source_text', 'target_text')})
        for change in ({'source_text':''}, {'target_text':'invented'},
                       {'target_text':'different'}, {'feedback_index':2}, {'feedback_index':False}):
            with self.subTest(change=change), self.assertRaises(ValueError):
                apply({**valid, **change})

    def test_few_shot_retained_or_replaced_once_without_mutating_old(self):
        previous = {'source_text':'old source', 'target_text':'old target'}
        old = [{'id':'a', 'text':'norm', 'weight':1., 'active':True, 'few_shot':previous}]
        feedback = [{'source_text':'new source', 'target_text':'new target'}]
        response = {'triggered_ids':['a'], 'new_rules':[]}
        self.assertEqual(learning.apply_memory_update(old, response, feedback)[0]['few_shot'], previous)
        update = {'id':'a', 'few_shot':{'feedback_index':0, **feedback[0]}}
        result = learning.apply_memory_update(old, {**response, 'few_shot_updates':[update]}, feedback)
        self.assertEqual(result[0]['few_shot'], feedback[0])
        self.assertEqual(old[0]['few_shot'], previous)
        for changes in ({'few_shot_updates':[update, update]},
                        {'triggered_ids':[], 'few_shot_updates':[update]}):
            with self.assertRaises(ValueError):
                learning.apply_memory_update(old, {**response, **changes}, feedback)

    def test_few_shot_excluded_from_quota_and_candidates_not_injected(self):
        example = {'source_text':'s' * 100, 'target_text':'t' * 100}
        old = [{'id':str(i), 'text':str(i) + 'n' * 499, 'weight':1., 'active':i < 6,
                'few_shot':example if i < 6 else {'source_text':'candidate example', 'target_text':'hidden'}}
               for i in range(7)]
        rules = learning.apply_memory_update(old, {'triggered_ids':['0'], 'new_rules':[]})
        self.assertEqual(sum(r['active'] for r in rules), 6)
        prompt = learning.memory_prompt({'rules':rules})
        self.assertEqual(prompt.count('Few shot'), 6)
        self.assertIn('s' * 100, prompt)
        self.assertNotIn('candidate example', prompt)

    def test_few_shot_worker_persistence_manual_edit_and_session_delete(self):
        self.feedback()
        reply = db.dumps_json({'triggered_ids':[], 'new_rules':[{'text':'norm', 'few_shot':{
            'feedback_index':0, 'source_text':'source0', 'target_text':'target0'}}]})
        with patch('term_extractor_app.ai_review.shared_provider.followup_chat', return_value=reply):
            learning._learning_worker(self.session)
        url = f'/api/ai-review/conversations/{self.session}/memory'
        memory = self.client.get(url).json()
        self.assertEqual(memory['rules'][0]['few_shot'], {'source_text':'source0', 'target_text':'target0'})
        self.assertEqual(learning.memory_snapshot(self.other)['rules'], [])
        edits = [{'id':memory['rules'][0]['id'], 'text':'norm', 'few_shot':memory['rules'][0]['few_shot']}]
        unchanged = self.client.put(url, json={'version':1, 'rules':edits}).json()
        self.assertIn('few_shot', unchanged['rules'][0])
        edits[0].update({'text':'rewritten norm', 'few_shot':{'source_text':'manual source', 'target_text':'manual target'}})
        changed = self.client.put(url, json={'version':1, 'rules':edits}).json()
        self.assertEqual(changed['rules'][0]['few_shot'], {'source_text':'manual source', 'target_text':'manual target'})
        self.assertEqual(changed['version'], 2)
        edits[0]['few_shot'] = {'source_text':'', 'target_text':''}
        deleted = self.client.put(url, json={'version':2, 'rules':edits}).json()
        self.assertNotIn('few_shot', deleted['rules'][0])
        self.assertEqual(deleted['version'], 3)
        delete_session(self.session)
        self.assertEqual(learning.memory_snapshot(self.session)['rules'], [])

    def test_manual_few_shot_requires_bilingual_pair(self):
        rules = [{'id':'a', 'text':'norm', 'weight':1., 'active':True,
                  'few_shot':{'source_text':'old', 'target_text':'旧'}}]
        with db.get_connection() as conn:
            conn.execute('INSERT INTO review_memory VALUES (?, ?, ?, ?)',
                         (self.session, 1, db.dumps_json(rules), 'now'))
        url = f'/api/ai-review/conversations/{self.session}/memory'
        for value in ({'source_text':'only source', 'target_text':''}, {'source_text':'', 'target_text':'仅译文'}):
            response = self.client.put(url, json={'version':1, 'rules':[{'id':'a', 'text':'norm', 'few_shot':value}]})
            self.assertEqual(response.status_code, 400)
        unchanged = self.client.put(url, json={'version':1, 'rules':[{'id':'a', 'text':'norm', 'few_shot':rules[0]['few_shot']}]}).json()
        self.assertEqual(unchanged['version'], 1)

    def test_manual_few_shot_can_be_added_to_an_old_rule(self):
        rules = [{'id':'a', 'text':'norm', 'weight':1., 'active':True}]
        with db.get_connection() as conn:
            conn.execute('INSERT INTO review_memory VALUES (?, ?, ?, ?)',
                         (self.session, 1, db.dumps_json(rules), 'now'))
        url = f'/api/ai-review/conversations/{self.session}/memory'
        response = self.client.put(url, json={'version':1, 'rules':[{'id':'a', 'text':'norm',
            'few_shot':{'source_text':'manual source', 'target_text':'manual target'}}]})
        self.assertEqual(response.status_code, 200)
        self.assertEqual(response.json()['rules'][0]['few_shot'],
                         {'source_text':'manual source', 'target_text':'manual target'})

    def test_invalid_case_keeps_memory_and_feedback_for_retry(self):
        self.feedback()
        reply = db.dumps_json({'triggered_ids':[], 'new_rules':[{'text':'norm', 'few_shot':{
            'feedback_index':0, 'source_text':'invented', 'target_text':'target0'}}]})
        with patch('term_extractor_app.ai_review.shared_provider.followup_chat', return_value=reply):
            learning._learning_worker(self.session)
        self.assertEqual(learning.memory_snapshot(self.session)['version'], 0)
        self.assertEqual(learning.learning_status(self.session)['events'][0]['status'], 'failed')
        self.assertFalse(review.get_review_results('t')[0]['has_issue'])

    def test_concurrent_manual_edit_not_overwritten_by_learning(self):
        with db.get_connection() as conn:
            conn.execute('INSERT INTO review_memory VALUES (?, ?, ?, ?)', (self.session, 1,
                         db.dumps_json([{'id':'a', 'text':'old', 'weight':1., 'active':True}]), 'now'))
        self.feedback()
        def respond(**kwargs):
            learning.update_memory(self.session, 1, [{'id':'a', 'text':'manual edit'}])
            return db.dumps_json({'triggered_ids':['a'], 'new_rules':[]})
        with patch('term_extractor_app.ai_review.shared_provider.followup_chat', side_effect=respond):
            learning._learning_worker(self.session)
        self.assertEqual(learning.memory_snapshot(self.session)['rules'][0]['text'], 'manual edit')
        self.assertEqual(learning.learning_status(self.session)['events'][0]['status'], 'failed')

    def test_memory_editor_api_preserves_weights_and_rejects_stale_save(self):
        rules = [
            {'id':'active','text':'使用中的规范','weight':2.5,'active':True},
            {'id':'candidate','text':'候补规范','weight':1.2,'active':False},
        ]
        with db.get_connection() as conn:
            conn.execute('INSERT INTO review_memory VALUES (?, ?, ?, ?)',
                         (self.session, 3, db.dumps_json(rules), 'now'))
        url=f'/api/ai-review/conversations/{self.session}/memory'
        loaded=self.client.get(url)
        self.assertEqual(loaded.status_code,200)
        self.assertEqual(loaded.json()['active_count'],1)
        self.assertEqual(loaded.json()['candidate_count'],1)
        payload={'version':3,'rules':[{'id':'active','text':'修改后的使用规范'},
                                      {'id':'candidate','text':'修改后的候补规范'}]}
        saved=self.client.put(url,json=payload)
        self.assertEqual(saved.status_code,200)
        self.assertEqual(saved.json()['version'],4)
        self.assertEqual([rule['weight'] for rule in saved.json()['rules']],[2.5,1.2])
        self.assertEqual([rule['active'] for rule in saved.json()['rules']],[True,False])
        stale=self.client.put(url,json=payload)
        self.assertEqual(stale.status_code,400)
        self.assertIn('重新打开',stale.json()['detail'])

    def test_memory_editor_rejects_missing_or_duplicate_rules(self):
        rules=[{'id':'a','text':'A','weight':1.0,'active':True},
               {'id':'b','text':'B','weight':1.0,'active':False}]
        with db.get_connection() as conn:
            conn.execute('INSERT INTO review_memory VALUES (?, ?, ?, ?)',
                         (self.session, 1, db.dumps_json(rules), 'now'))
        url=f'/api/ai-review/conversations/{self.session}/memory'
        missing=self.client.put(url,json={'version':1,'rules':[{'id':'a','text':'A'}]})
        self.assertEqual(missing.status_code,400)
        duplicate=self.client.put(url,json={'version':1,'rules':[{'id':'a','text':'same'},{'id':'b','text':'same'}]})
        self.assertEqual(duplicate.status_code,400)

    def test_learning_retry_applies_weight_once(self):
        event=self.feedback()
        with patch('term_extractor_app.ai_review.shared_provider.followup_chat',return_value='invalid'):
            learning._learning_worker(self.session)
        self.assertEqual(learning.memory_snapshot(self.session)['version'],0)
        learning.retry_learning(self.session,event['event_id'])
        with patch('term_extractor_app.ai_review.shared_provider.followup_chat',return_value='{"triggered_ids":[],"new_rules":["规范"]}'):
            learning._learning_worker(self.session)
        learning.retry_learning(self.session,event['event_id'])
        with patch('term_extractor_app.ai_review.shared_provider.followup_chat') as call:
            learning._learning_worker(self.session)
            call.assert_not_called()
        self.assertEqual(learning.memory_snapshot(self.session)['version'],1)

    def test_same_session_workers_are_serial(self):
        self.feedback(0); self.feedback(1)
        calls=[]
        def respond(**kwargs):
            payload=db.loads_json(kwargs['messages'][-1]['content'],{})
            calls.append(payload['operation'])
            time.sleep(.05)
            old=payload['old_rules']
            return db.dumps_json({'triggered_ids':[old[0]['id']] if old else [], 'new_rules':[] if old else ['参考规范']})
        self.patches[2].stop()
        with patch('term_extractor_app.ai_review.shared_provider.followup_chat',side_effect=respond):
            learning.start_learning(self.session); learning.start_learning(self.session)
            deadline=time.monotonic()+5
            while self.session in learning._workers and time.monotonic()<deadline:
                time.sleep(.01)
            self.assertNotIn(self.session, learning._workers)
        self.assertEqual(calls,['create','update'])
        self.assertEqual(learning.memory_snapshot(self.session)['rules'][0]['weight'],2)

    def test_cache_no_issue_only_isolation_and_toggle(self):
        learning.cache_results(self.session, [(self.items[0], {'has_issue':False}), (self.items[1], {'has_issue':True}), (self.items[2], {})])
        self.assertEqual(set(learning.cached_results(self.session, ['k0','k1','k2'])), {'k0'})
        self.assertEqual(learning.cached_results(self.other, ['k0']), {})
        learning.save_options({'cache_enabled':False})
        learning.cache_results(self.session, [(self.items[1], {'has_issue':False})])
        self.assertEqual(learning.cached_results(self.session, ['k0']), {})
        learning.save_options({'cache_enabled':True})
        self.assertEqual(set(learning.cached_results(self.session, ['k0','k1'])), {'k0'})

    def test_feedback_preserves_original_updates_cache_and_is_idempotent(self):
        first = self.feedback()
        self.assertTrue(first['event_id'])
        self.assertFalse(review.get_review_results('t')[0]['has_issue'])
        self.assertFalse(learning.cached_results(self.session, ['k0'])['k0']['has_issue'])
        self.assertEqual(self.feedback()['changed'], 0)
        with db.get_connection() as conn:
            row = conn.execute('SELECT original_json FROM review_feedback').fetchone()
            self.assertTrue(db.loads_json(row[0], {})['has_issue'])
            self.assertEqual(conn.execute('SELECT COUNT(*) FROM review_learning_events').fetchone()[0], 1)

    def test_learning_off_preserves_feedback(self):
        learning.save_options({'learning_enabled':False})
        self.assertFalse(self.feedback()['event_id'])
        self.assertFalse(review.get_review_results('t')[0]['has_issue'])

    def test_learning_request_contains_only_disagree_and_marks_rejected_advice(self):
        learning.submit_feedback(self.session, 't', [
            {'result_id':self.results[0]['id'], 'decision':'disagree', 'reason':'允许自然改写'},
            {'result_id':self.results[1]['id'], 'decision':'agree', 'reason':'确实有错'},
        ])
        with patch('term_extractor_app.ai_review.shared_provider.followup_chat',
                   return_value='{"triggered_ids":[],"new_rules":["结合文体判断自然改写"]}') as call:
            learning._learning_worker(self.session)
        payload=db.loads_json(call.call_args.kwargs['messages'][-1]['content'], {})
        self.assertEqual(len(payload['feedback']),1)
        evidence=payload['feedback'][0]
        self.assertEqual(evidence['result_id'],self.results[0]['id'])
        self.assertEqual(evidence['source_text'],'source0')
        self.assertEqual(evidence['user_reason'],'允许自然改写')
        self.assertEqual(evidence['user_decision'],'disagree')
        self.assertNotIn('suggestion', evidence)
        self.assertEqual(evidence['rejected_ai_judgment']['suggestion'],'suggestion')
        self.assertTrue(review.get_review_results('t')[1]['has_issue'])
        self.assertTrue(review.get_review_results('t')[2]['has_issue'])

    def test_learning_prompt_handles_historical_event_without_mutating_evidence(self):
        item={'source_text':'战旗', 'target_text':'Tactical Strategy',
              'reason':'战旗是游戏类型', 'suggestion':'War Banner',
              'term_reference':[{'source':'term','targets':['target']}]}
        messages=learning.build_learning_messages({'version':0,'rules':[]},[item])
        payload=db.loads_json(messages[-1]['content'],{})
        self.assertEqual(payload['feedback'][0]['rejected_ai_judgment']['suggestion'],'War Banner')
        self.assertEqual(payload['feedback'][0]['user_reason'],'战旗是游戏类型')
        self.assertIn('term_reference', item)

    def test_feedback_atomic_rollback(self):
        entries=[{'result_id':self.results[0]['id'],'decision':'disagree'}, {'result_id':'missing','decision':'disagree'}]
        with self.assertRaises(ValueError): learning.submit_feedback(self.session,'t',entries)
        self.assertTrue(review.get_review_results('t')[0]['has_issue'])
        self.assertEqual(learning.cached_results(self.session,['k0']), {})

    def test_excel_round_trip_sort_validation_style_and_duplicate(self):
        path = output.generate_review_excel('t')
        book = load_workbook(path)
        sheet=book['审校结果']
        headers=[c.value for c in sheet[1]]
        self.assertEqual(headers[-3:], ['agree/disagree','reason','_result_id'])
        self.assertEqual(book['_review_meta'].sheet_state, 'veryHidden')
        self.assertTrue(sheet.column_dimensions['L'].hidden)
        self.assertEqual(len(sheet.data_validations.dataValidation),1)
        self.assertEqual(len(sheet.conditional_formatting),1)
        a=[c.value for c in sheet[2]]; b=[c.value for c in sheet[4]]
        for col, value in enumerate(b,1): sheet.cell(2,col,value)
        for col, value in enumerate(a,1): sheet.cell(4,col,value)
        sheet.cell(2,10,'disagree'); sheet.cell(2,11,'自然表达')
        buffer=io.BytesIO(); book.save(buffer); book.close()
        result=learning.import_feedback(self.session, buffer.getvalue())
        self.assertEqual(result['changed'],1)
        self.assertFalse(review.get_review_results('t')[2]['has_issue'])
        self.assertTrue(review.get_review_results('t')[0]['has_issue'])
        self.assertEqual(learning.import_feedback(self.session, buffer.getvalue())['changed'],0)
        with self.assertRaises(ValueError): learning.import_feedback(self.other, buffer.getvalue())
        self.assertEqual(len(output.NORMAL_HEADERS),9)

    def test_excel_rejects_modified_source_and_old_export(self):
        book=load_workbook(output.generate_review_excel('t')); sheet=book['审校结果']
        sheet.cell(2,4,'tampered'); sheet.cell(2,10,'disagree')
        buf=io.BytesIO(); book.save(buf)
        with self.assertRaises(ValueError): learning.import_feedback(self.session,buf.getvalue())
        del book['_review_meta']; buf=io.BytesIO(); book.save(buf); book.close()
        with self.assertRaises(ValueError): learning.import_feedback(self.session,buf.getvalue())

    def test_agree_no_learning_and_formula_safe_export(self):
        result=learning.submit_feedback(self.session,'t',[{'result_id':self.results[0]['id'],'decision':'agree','reason':'=HYPERLINK("bad")'}])
        self.assertFalse(result['event_id'])
        book=load_workbook(result['output_path']); self.assertEqual(book['审校结果']['K2'].data_type,'s'); book.close()
        self.assertTrue(review.get_review_results('t')[0]['has_issue'])

    def test_export_failure_does_not_undo_feedback(self):
        with patch.object(output,'generate_review_excel',side_effect=PermissionError('locked')):
            result=self.feedback()
        self.assertIn('locked',result['export_error'])
        self.assertFalse(review.get_review_results('t')[0]['has_issue'])
        self.assertFalse(learning.regenerate_output(self.session,'t')['export_error'])

    def test_weights_candidate_swap_and_duplicate_trigger(self):
        old=[{'id':str(i),'text':'规'*500,'weight':1.,'active':i<6} for i in range(7)]
        rules=learning.apply_memory_update(old, {'triggered_ids':['6','6'],'new_rules':[]})
        self.assertEqual(rules[6]['weight'],2)
        self.assertEqual(rules[0]['weight'],.9)
        self.assertTrue(rules[6]['active']); self.assertFalse(rules[5]['active'])
        self.assertEqual(sum(r['active'] for r in rules),6)

    def test_learning_create_update_and_failure_preserves_old(self):
        first=self.feedback()
        with patch('term_extractor_app.ai_review.shared_provider.followup_chat', return_value='{"triggered_ids":[],"new_rules":["英语中自然的名词表达不机械替换为动词"]}') as call:
            learning._learning_worker(self.session)
            self.assertEqual(call.call_count,1)
        memory=learning.memory_snapshot(self.session)
        self.assertEqual(memory['version'],1)
        self.feedback(1)
        reply=db.dumps_json({'triggered_ids':[memory['rules'][0]['id']],'new_rules':[]})
        with patch('term_extractor_app.ai_review.shared_provider.followup_chat',return_value=reply): learning._learning_worker(self.session)
        self.assertEqual(learning.memory_snapshot(self.session)['rules'][0]['weight'],2)
        self.feedback(2)
        with patch('term_extractor_app.ai_review.shared_provider.followup_chat',side_effect=TimeoutError('timeout')): learning._learning_worker(self.session)
        self.assertEqual(learning.memory_snapshot(self.session)['version'],2)
        self.assertEqual(learning.learning_status(self.session)['events'][0]['status'],'failed')
        self.assertFalse(review.get_review_results('t')[2]['has_issue'])

    def test_deleted_session_late_worker_cannot_recreate_memory_or_cache(self):
        self.feedback()
        def late(**kwargs):
            delete_session(self.session)
            return '{"triggered_ids":[],"new_rules":["参考规范"]}'
        with patch('term_extractor_app.ai_review.shared_provider.followup_chat',side_effect=late): learning._learning_worker(self.session)
        learning.cache_results(self.session, [(self.items[0],{'has_issue':False})])
        self.assertEqual(learning.memory_snapshot(self.session)['version'],0)
        self.assertEqual(learning.cached_results(self.session,['k0']),{})

    def test_workflow_batches_only_misses_and_memory_does_not_change_cache(self):
        with db.get_connection() as conn:
            conn.execute('DELETE FROM review_results')
        packages=[]
        async def run(**kwargs):
            packages.extend(kwargs['packages'])
            for package in kwargs['packages']:
                pairs=[(item, {'id':item['id'],'has_issue':False,'issue':'','issue_type':'','suggestion':''}) for item in package]
                review._save_review_results_bulk('t',pairs)
                learning.cache_results(self.session,pairs)
            return len(kwargs['packages']),3,0
        with patch.object(review,'get_shared_ai_settings',return_value={}), patch.object(review,'_run_review_packages',side_effect=run), patch.object(review,'_track_ai_review_finish'):
            review._run_review_task_impl('t')
            self.assertEqual(sum(len(p) for p in packages),3)
            packages.clear()
            with db.get_connection() as conn:
                conn.execute('DELETE FROM review_results')
                config={**self.config,'memory_snapshot':{'version':2,'rules':[{'text':'new memory','active':True}]}}
                conn.execute('UPDATE review_tasks SET config_json=? WHERE id=?',(db.dumps_json(config),'t'))
            review._run_review_task_impl('t')
            self.assertEqual(packages,[])
            self.assertEqual(review.get_review_task('t')['cached_count'],3)
            with db.get_connection() as conn:
                conn.execute('DELETE FROM review_results')
                conn.execute("UPDATE file_items SET target_text='changed' WHERE id='i1'")
            review._run_review_task_impl('t')
            self.assertEqual(sum(len(p) for p in packages),1)


if __name__ == '__main__':
    unittest.main()
