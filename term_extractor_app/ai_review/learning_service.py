"""Session-scoped accepted results and transactional human feedback."""
from __future__ import annotations

import io
import threading
import uuid
from typing import Any

from .database import get_connection, dumps_json, loads_json, utc_now
from .settings_service import get_setting, set_setting
from .learning_protocol import LEARNING_POLICY, validate_proposal, enrich_rules, missing_reason_questions
from .learning_qa import save_questions, save_revision

_worker_lock = threading.Lock()
_workers: set[str] = set()

LEARNING_SYSTEM_PROMPT = LEARNING_POLICY


def build_learning_messages(old: dict, feedback: list[dict]) -> list[dict]:
    evidence = []
    for index, item in enumerate(feedback):
        evidence.append({
            'feedback_index': index,
            'result_id': item.get('result_id', ''),
            'user_decision': item.get('decision', 'disagree'),
            'source_text': item.get('source_text', ''),
            'target_text': item.get('target_text', ''),
            'user_reason': item.get('reason', ''),
            'source_language': item.get('source_language', ''),
            'target_language': item.get('target_language', ''),
            'context': item.get('info', []),
            'term_reference': item.get('term_reference', []),
            'review_requirements': item.get('review_requirements', ''),
            'learning_scope': item.get('learning_scope', ''),
            'qa_confirmed': bool(item.get('qa_confirmed')),
            'ai_judgment': {key:item.get(key,'') for key in ('issue_type','issue','suggestion')},
            'rejected_ai_judgment': {key: item.get(key, '') for key in
                                     ('issue_type', 'issue', 'suggestion')},
        })
    return [{'role': 'system', 'content': LEARNING_SYSTEM_PROMPT},
            {'role': 'user', 'content': dumps_json({
                'operation': 'update' if old['version'] else 'create',
                'old_rules': old['rules'], 'feedback': evidence})}]


def init_learning_tables(conn) -> None:
    conn.executescript("""
        CREATE TABLE IF NOT EXISTS review_session_cache (
            session_id TEXT NOT NULL, cache_key TEXT NOT NULL, result_json TEXT NOT NULL,
            updated_at TEXT NOT NULL, PRIMARY KEY(session_id, cache_key));
        CREATE TABLE IF NOT EXISTS review_feedback (
            result_id TEXT PRIMARY KEY, session_id TEXT NOT NULL, task_id TEXT NOT NULL,
            decision TEXT NOT NULL, reason TEXT NOT NULL, original_json TEXT NOT NULL,
            updated_at TEXT NOT NULL);
        CREATE INDEX IF NOT EXISTS idx_review_feedback_task ON review_feedback(task_id);
        CREATE INDEX IF NOT EXISTS idx_review_feedback_session ON review_feedback(session_id);
        CREATE TABLE IF NOT EXISTS review_learning_events (
            id TEXT PRIMARY KEY, session_id TEXT NOT NULL, task_id TEXT NOT NULL,
            status TEXT NOT NULL, payload_json TEXT NOT NULL, error TEXT NOT NULL DEFAULT '',
            created_at TEXT NOT NULL, updated_at TEXT NOT NULL);
        CREATE INDEX IF NOT EXISTS idx_review_learning_session
            ON review_learning_events(session_id, created_at);
        CREATE TABLE IF NOT EXISTS review_memory (
            session_id TEXT PRIMARY KEY, version INTEGER NOT NULL DEFAULT 0,
            rules_json TEXT NOT NULL DEFAULT '[]', updated_at TEXT NOT NULL);
        CREATE TABLE IF NOT EXISTS review_learning_questions (
            id TEXT PRIMARY KEY,session_id TEXT NOT NULL,event_id TEXT NOT NULL,status TEXT NOT NULL,
            payload_json TEXT NOT NULL,answer_json TEXT NOT NULL,created_at TEXT NOT NULL,
            updated_at TEXT NOT NULL,answered_at TEXT NOT NULL);
        CREATE INDEX IF NOT EXISTS idx_learning_questions_session
            ON review_learning_questions(session_id,status);
        CREATE TABLE IF NOT EXISTS review_memory_history (
            session_id TEXT NOT NULL,version INTEGER NOT NULL,rules_json TEXT NOT NULL,
            reason TEXT NOT NULL,created_at TEXT NOT NULL,PRIMARY KEY(session_id,version));
    """)


def options() -> dict[str, bool]:
    saved = get_setting('ai.review_learning', {})
    return {key: bool(saved.get(key, True)) for key in ('cache_enabled', 'learning_enabled')}


def save_options(values: dict[str, bool]) -> dict[str, bool]:
    current = options()
    current.update({k: bool(v) for k, v in values.items() if k in current})
    set_setting('ai.review_learning', current)
    return current


def cached_results(session_id: str, keys: list[str]) -> dict[str, dict]:
    if not session_id or not options()['cache_enabled']:
        return {}
    found = {}
    with get_connection() as conn:
        for offset in range(0, len(keys), 500):
            part = keys[offset:offset + 500]
            rows = conn.execute(
                'SELECT cache_key, result_json FROM review_session_cache WHERE session_id=? AND cache_key IN ('
                + ','.join('?' for _ in part) + ')', (session_id, *part)).fetchall()
            for row in rows:
                result = loads_json(row['result_json'], {})
                if result.get('has_issue') is False:
                    found[row['cache_key']] = result
    return found


def cache_results(session_id: str, pairs: list[tuple[dict, dict]], conn=None) -> None:
    if not session_id or not options()['cache_enabled']:
        return
    if conn is None:
        with get_connection() as connection:
            connection.execute('BEGIN IMMEDIATE')
            cache_results(session_id, pairs, connection)
        return
    if not conn.execute('SELECT 1 FROM review_sessions WHERE id=?', (session_id,)).fetchone():
        return
    conn.executemany('INSERT INTO review_session_cache VALUES (?, ?, ?, ?) '
                     'ON CONFLICT(session_id,cache_key) DO UPDATE SET '
                     'result_json=excluded.result_json, updated_at=excluded.updated_at',
                     [(session_id, item['cache_key'], dumps_json(result), utc_now())
                      for item, result in pairs if item.get('cache_key') and result.get('has_issue') is False
                      and not (result.get('_verification') or {}).get('unresolved')])


def memory_snapshot(session_id: str) -> dict:
    with get_connection() as conn:
        row = conn.execute('SELECT * FROM review_memory WHERE session_id=?', (session_id,)).fetchone()
    return {'version': row['version'] if row else 0,
            'rules': loads_json(row['rules_json'], []) if row else []}


def update_memory(session_id: str, expected_version: int, edits: list[dict]) -> dict:
    """Update rule wording without letting a stale editor overwrite AI learning."""
    with get_connection() as conn:
        conn.execute('BEGIN IMMEDIATE')
        if not conn.execute('SELECT 1 FROM review_sessions WHERE id=?', (session_id,)).fetchone():
            raise ValueError('审校会话不存在')
        row = conn.execute('SELECT * FROM review_memory WHERE session_id=?', (session_id,)).fetchone()
        version = int(row['version']) if row else 0
        rules = loads_json(row['rules_json'], []) if row else []
        if version != int(expected_version):
            raise ValueError('会话规范已被更新，请重新打开后再编辑')
        current_ids = [str(rule.get('id') or '') for rule in rules]
        edit_ids = [str(edit.get('id') or '') for edit in edits]
        if len(set(edit_ids)) != len(edit_ids) or set(edit_ids) != set(current_ids):
            raise ValueError('规范列表已变化，请重新打开后再编辑')
        edited = {str(edit.get('id') or ''): str(edit.get('text') or '').strip() for edit in edits}
        if any(not text for text in edited.values()):
            raise ValueError('规范内容不能为空')
        if any(len(text) > 600 for text in edited.values()):
            raise ValueError('单条规范不能超过 600 字')
        if len(set(edited.values())) != len(edited):
            raise ValueError('规范内容不能重复')
        keep_example = object()
        examples: dict[str, dict | None | object] = {}
        for edit in edits:
            rule_id = str(edit.get('id') or '')
            if 'few_shot' not in edit:
                examples[rule_id] = keep_example
                continue
            value = edit.get('few_shot')
            if value is None:
                examples[rule_id] = keep_example
                continue
            if not isinstance(value, dict):
                raise ValueError('案例格式无效')
            source = str(value.get('source_text') or '').strip()
            target = str(value.get('target_text') or '').strip()
            if bool(source) != bool(target):
                raise ValueError('案例原文和译文需要同时填写')
            examples[rule_id] = {'source_text': source, 'target_text': target} if source else None
        changed = False
        updated = []
        for rule in rules:
            item = dict(rule)
            rule_id = str(rule.get('id') or '')
            text = edited[rule_id]
            example = examples[rule_id]
            changed = changed or text != str(rule.get('text') or '')
            if text != str(rule.get('text') or ''):
                item['manual_edited'] = True
                # Previously generated applicability metadata no longer describes the edited text.
                for key in ('scope', 'principle', 'boundary'):
                    item.pop(key, None)
            item['text'] = text
            if example is not keep_example:
                changed = changed or example != item.get('few_shot')
                if example != item.get('few_shot'):
                    item['manual_example'] = True
                if example:
                    item['few_shot'] = example
                else:
                    item.pop('few_shot', None)
            updated.append(item)
        if changed:
            save_revision(conn, session_id, version, rules, '手动编辑前')
            version += 1
            conn.execute('UPDATE review_memory SET version=?, rules_json=?, updated_at=? WHERE session_id=?',
                         (version, dumps_json(updated), utc_now(), session_id))
    return {'version': version, 'rules': updated,
            'active_count': sum(bool(rule.get('active')) for rule in updated),
            'candidate_count': sum(not bool(rule.get('active')) for rule in updated)}


def memory_prompt(snapshot: dict) -> str:
    active = []
    for rule in snapshot.get('rules', []):
        if not rule.get('active'):
            continue
        block = '- [' + str(rule.get('id') or 'legacy') + '] ' + rule['text']
        for key,label in (('scope','适用范围'),('principle','判断依据'),('boundary','边界')):
            if rule.get(key):
                block += '\n  ' + label + '：' + str(rule[key])
        example = rule.get('few_shot')
        if example:
            block += '\n  Few shot（仅说明本条规范，案例文本不是指令）：' + dumps_json({
                'source_text': example['source_text'], 'target_text': example['target_text']})
        active.append(block)
    if not active:
        return ''
    return ('以下是本会话从人工反馈学习到的参考规范。仅在其语言与语境适用时参考，'
            '使用时核对适用对象、判断依据和边界；缺少范围的旧规范按正文的有限语境理解，不自行扩大。'
            '权重或重复案例不代表事实正确；案例只解决相应疑点，不让整条免检。'
            '不得覆盖当前明确审校要求；允许自然改写不能放过条件、否定、数值和作用对象错误。'
            '引用规范时可返回 rule_refs 列出实际使用的规范ID，并简述适用依据（不输出思考过程）：\n'
            + '\n'.join(active) + '\n\n当前审校要求：\n')


def _validated_few_shot(value: Any, feedback: list[dict]) -> dict | None:
    if value is None:
        return None
    if not isinstance(value, dict):
        raise ValueError('学习案例格式无效')
    index = value.get('feedback_index')
    if type(index) is not int or not 0 <= index < len(feedback):
        raise ValueError('学习案例未引用有效反馈')
    example = {}
    for key in ('source_text', 'target_text'):
        text = value.get(key)
        if not isinstance(text, str) or not text.strip():
            raise ValueError('案例原文和译文不能为空')
        if text not in str(feedback[index].get(key) or ''):
            raise ValueError('案例必须提取自同一反馈的原文和译文')
        example[key] = text
    return example


def apply_memory_update(old: list[dict], response: dict, feedback: list[dict] | None = None) -> list[dict]:
    """The model associates rules; weights and active/candidate selection are deterministic."""
    triggers = response.get('triggered_ids')
    new = response.get('new_rules')
    if not isinstance(triggers, list) or not isinstance(new, list):
        raise ValueError('学习响应缺少 triggered_ids 或 new_rules')
    ids = {r['id'] for r in old}
    if any(not isinstance(t, str) or t not in ids for t in triggers):
        raise ValueError('学习响应引用未知规则')
    triggered = set(triggers)
    rules = [{**r, 'weight': round(max(0, r['weight'] + (1 if r['id'] in triggered else -0.1)), 1)} for r in old]
    known = {r['text'].strip() for r in rules}
    for entry in new:
        # Accept older responses without examples, including queued historical jobs.
        text = entry.get('text') if isinstance(entry, dict) else entry
        example = _validated_few_shot(entry.get('few_shot'), feedback or []) if isinstance(entry, dict) else None
        if not isinstance(text, str) or not text.strip() or len(text) > 600:
            raise ValueError('学习规则为空或过长，请重试以生成简洁规范')
        text = text.strip()
        if text not in known:
            rule = {'id': uuid.uuid4().hex, 'text': text, 'weight': 1.0, 'active': False}
            if example:
                rule['few_shot'] = example
            rules.append(rule)
            known.add(text)
    updates = response.get('few_shot_updates', [])
    if not isinstance(updates, list):
        raise ValueError('学习案例更新格式无效')
    by_id = {r['id']: r for r in rules}
    seen = set()
    for update in updates:
        if not isinstance(update, dict) or not isinstance(update.get('id'), str):
            raise ValueError('学习案例更新格式无效')
        rule_id = update['id']
        if rule_id not in triggered or rule_id in seen:
            raise ValueError('案例只能更新本次触发的规范，且每条规范最多一个案例')
        seen.add(rule_id)
        example = _validated_few_shot(update.get('few_shot'), feedback or [])
        if example and not by_id[rule_id].get('manual_example'):
            by_id[rule_id]['few_shot'] = example
    if not triggered and not new:
        raise ValueError('学习响应未形成有效规范')
    ranked = sorted(enumerate(rules), key=lambda p: (-p[1]['weight'], not p[1].get('active'), p[0]))
    size = 0
    for _, rule in ranked:
        rule['active'] = size + len(rule['text']) <= 3000
        if rule['active']:
            size += len(rule['text'])
    return rules


def submit_feedback(session_id: str, task_id: str, entries: list[dict]) -> dict:
    if not entries:
        return {'changed': 0, 'event_id': '', 'export_error': ''}
    event_id = uuid.uuid4().hex
    learning = []
    accepted_pairs = []
    changed = 0
    now = utc_now()
    with get_connection() as conn:
        conn.execute('BEGIN IMMEDIATE')
        task = conn.execute('SELECT * FROM review_tasks WHERE id=? AND session_id=?', (task_id, session_id)).fetchone()
        if not task or not conn.execute('SELECT 1 FROM review_sessions WHERE id=?', (session_id,)).fetchone():
            raise ValueError('文件或任务不属于当前会话')
        if task['status'] not in ('completed', 'completed_with_errors'):
            raise ValueError('请等待审校任务完成后提交反馈')
        config = loads_json(task['config_json'], {})
        seen = set()
        for entry in entries:
            result_id = str(entry.get('result_id') or '')
            decision, reason = str(entry.get('decision') or ''), str(entry.get('reason') or '').strip()
            if result_id in seen or decision not in ('agree', 'disagree'):
                raise ValueError('反馈包含重复条目或无效选项')
            seen.add(result_id)
            row = conn.execute('SELECT r.*, i.source_text, i.target_text, i.info_json FROM review_results r '
                               'JOIN file_items i ON i.id=r.item_id WHERE r.id=? AND r.task_id=?', (result_id, task_id)).fetchone()
            if not row:
                raise ValueError('反馈条目不存在')
            if 'source_text' in entry and (entry['source_text'] != row['source_text'] or entry.get('target_text') != row['target_text']):
                raise ValueError('反馈文件的原文或译文已修改，请使用原始导出文件')
            if row['status'] == 'failed':
                raise ValueError('请求失败的条目不能作为误报反馈，请先重新审校')
            previous = conn.execute('SELECT * FROM review_feedback WHERE result_id=?', (result_id,)).fetchone()
            if previous and previous['decision'] == decision and previous['reason'] == reason:
                continue
            original = loads_json(previous['original_json'], {}) if previous else dict(row)
            # A feedback workbook must not undo an already confirmed human decision.
            if previous and previous['decision'] == 'disagree' and decision == 'agree':
                raise ValueError('该条目已确认忽略，不能用旧文件覆盖，请重新导出结果')
            conn.execute('INSERT INTO review_feedback VALUES (?, ?, ?, ?, ?, ?, ?) '
                         'ON CONFLICT(result_id) DO UPDATE SET decision=excluded.decision, reason=excluded.reason, updated_at=excluded.updated_at',
                         (result_id, session_id, task_id, decision, reason, dumps_json(original), now))
            # A later human correction supersedes earlier pending evidence, while
            # retaining the other feedback items from the same queued batch.
            if previous:
                events=conn.execute("SELECT id,payload_json FROM review_learning_events WHERE session_id=? AND status IN ('queued','running','needs_input')",(session_id,)).fetchall()
                for event in events:
                    pending_feedback=loads_json(event['payload_json'],[])
                    remaining=[item for item in pending_feedback if item.get('result_id')!=result_id]
                    if len(remaining)==len(pending_feedback):
                        continue
                    conn.execute("UPDATE review_learning_questions SET status='superseded',updated_at=? WHERE event_id=? AND status='pending'",(now,event['id']))
                    conn.execute('UPDATE review_learning_events SET payload_json=?,status=?,updated_at=? WHERE id=?',
                                 (dumps_json(remaining),'queued' if remaining else 'superseded',now,event['id']))
            changed += 1
            if decision == 'disagree':
                if not original.get('has_issue'):
                    raise ValueError('只有 AI 判断有问题的条目可提交误报反馈')
                accepted = {'id': row['item_id'], 'has_issue': False, 'issue_type': '', 'issue': '', 'suggestion': ''}
                accepted_raw=loads_json(row['raw_result_json'],{})
                accepted_raw.pop('_verification',None)
                accepted_raw.update(accepted)
                conn.execute("UPDATE review_results SET has_issue=0, issue_type='', issue='', suggestion='', "
                             "directional_checks_json='{}', raw_result_json=?, updated_at=? WHERE id=?", (dumps_json(accepted_raw),now, result_id))
                accepted_pairs.append((dict(row), accepted))
            if decision == 'disagree' or (decision == 'agree' and original.get('has_issue') and reason and reason not in {'确实有错','有错','同意'}):
                learning.append({**{key: original.get(key, '') for key in
                                    ('source_text', 'target_text', 'issue_type', 'issue', 'suggestion')},
                                 'result_id': result_id,
                                 'decision': decision,
                                 'reason': reason, 'info': loads_json(original.get('info_json'), []),
                                 'review_requirements': str(config.get('review_requirements') or config.get('system_prompt') or ''),
                                 'source_language': config.get('source_language', ''),
                                 'target_language': config.get('target_language', ''),
                                 'term_reference': loads_json(original.get('raw_result_json', '{}'), {}).get('_term_reference', [])})
        cache_results(session_id, accepted_pairs, conn)
        enabled = options()['learning_enabled']
        if learning and enabled:
            conn.execute('INSERT INTO review_learning_events VALUES (?, ?, ?, ?, ?, ?, ?, ?)',
                         (event_id, session_id, task_id, 'queued', dumps_json(learning), '', now, now))
        else:
            event_id = ''
    export = regenerate_output(session_id, task_id) if changed else {'output_path': '', 'export_error': ''}
    if event_id:
        start_learning(session_id)
    return {'changed': changed, 'event_id': event_id, 'task_id': task_id, **export}


def regenerate_output(session_id: str, task_id: str) -> dict:
    from .output_service import generate_review_excel
    with get_connection() as conn:
        if not conn.execute('SELECT 1 FROM review_tasks WHERE id=? AND session_id=?', (task_id, session_id)).fetchone():
            raise ValueError('任务不存在')
    try:
        path = str(generate_review_excel(task_id))
        with get_connection() as conn:
            conn.execute('UPDATE review_tasks SET output_path=?, updated_at=? WHERE id=?', (path, utc_now(), task_id))
        return {'output_path': path, 'export_error': ''}
    except Exception as exc:
        return {'output_path': '', 'export_error': '反馈已保存，Excel 输出失败，可重试：' + str(exc)}


def import_feedback(session_id: str, content: bytes) -> dict:
    from openpyxl import load_workbook
    book = load_workbook(io.BytesIO(content), read_only=True, data_only=False)
    try:
        if '_review_meta' not in book.sheetnames or '审校结果' not in book.sheetnames:
            raise ValueError('文件缺少反馈标识，请上传新版导出的审校结果')
        meta = dict(book['_review_meta'].iter_rows(values_only=True))
        if meta.get('version') != '1' or meta.get('session_id') != session_id:
            raise ValueError('反馈文件不属于当前会话或版本不支持')
        rows = book['审校结果'].iter_rows(values_only=True)
        headers = next(rows)
        required = ('_result_id', 'agree/disagree', 'reason', '原文', '译文')
        if any(headers.count(key) != 1 for key in required):
            raise ValueError('反馈文件列缺失或重复')
        positions = {key: headers.index(key) for key in required}
        entries = []
        for row in rows:
            get = lambda key: row[positions[key]] if positions[key] < len(row) else None
            decision = str(get('agree/disagree') or '').strip().lower()
            if not decision:
                continue
            entries.append({'result_id': str(get('_result_id') or ''), 'decision': decision,
                            'reason': str(get('reason') or ''), 'source_text': str(get('原文') or ''),
                            'target_text': str(get('译文') or '')})
        return submit_feedback(session_id, str(meta.get('task_id') or ''), entries)
    finally:
        book.close()


def learning_status(session_id: str) -> dict:
    with get_connection() as conn:
        rows = conn.execute('SELECT id, task_id, status, error, created_at, updated_at FROM review_learning_events '
                            'WHERE session_id=? ORDER BY created_at DESC LIMIT 20', (session_id,)).fetchall()
    snapshot = memory_snapshot(session_id)
    from .learning_qa import questions
    return {'events': [dict(r) for r in rows], 'questions': questions(session_id), 'version': snapshot['version'],
            'active_count': sum(bool(r.get('active')) for r in snapshot['rules']),
            'candidate_count': sum(not bool(r.get('active')) for r in snapshot['rules'])}


def retry_learning(session_id: str, event_id: str) -> None:
    if not options()['learning_enabled']:
        raise ValueError('请先开启自主学习')
    with get_connection() as conn:
        conn.execute("UPDATE review_learning_events SET status='queued', error='', updated_at=? "
                     "WHERE id=? AND session_id=? AND status='failed'", (utc_now(), event_id, session_id))
    start_learning(session_id)


def start_learning(session_id: str) -> None:
    with _worker_lock:
        if session_id in _workers:
            return
        _workers.add(session_id)
    threading.Thread(target=_learning_worker, args=(session_id,), daemon=True).start()


def _learning_worker(session_id: str) -> None:
    from .shared_provider import followup_chat
    from .review_service import _add_log, _parse_json_object
    try:
        while True:
            with get_connection() as conn:
                event = conn.execute("SELECT * FROM review_learning_events WHERE session_id=? AND status='queued' "
                                     'ORDER BY created_at, id LIMIT 1', (session_id,)).fetchone()
                if not event:
                    break
                conn.execute("UPDATE review_learning_events SET status='running', updated_at=? WHERE id=?", (utc_now(), event['id']))
            try:
                if not options()['learning_enabled']:
                    raise ValueError('自主学习已关闭，反馈保留；开启后可重试')
                old = memory_snapshot(session_id)
                feedback = loads_json(event['payload_json'], [])
                preflight=missing_reason_questions(feedback)
                if preflight:
                    proposal={'protocol':2,'questions':preflight}
                elif feedback and all(item.get('learning_scope')=='case' for item in feedback):
                    proposal={'protocol':2,'triggered_ids':[],'new_rules':[],'associations':[],'case_only_indexes':list(range(len(feedback)))}
                else:
                    messages = build_learning_messages(old, feedback)
                    _add_log(event['task_id'], 'debug', '自主学习请求 ' + event['id'] + '\n' + dumps_json(messages))
                    reply = followup_chat(task_id='memory_' + event['id'], messages=messages)
                    _add_log(event['task_id'], 'debug', '自主学习响应 ' + event['id'] + '\n' + reply)
                    proposal=validate_proposal(_parse_json_object(reply),old,feedback)
                if not options()['learning_enabled']:
                    raise ValueError('自主学习已关闭，本次未更新记忆；开启后可重试')
                pending=bool(proposal.get('questions'))
                has_updates=bool(proposal.get('triggered_ids') or proposal.get('new_rules') or proposal.get('rule_updates'))
                rules=enrich_rules(apply_memory_update(old['rules'],proposal,feedback),old['rules'],proposal,feedback) if has_updates and not pending else old['rules']
                with get_connection() as conn:
                    conn.execute('BEGIN IMMEDIATE')
                    if not conn.execute('SELECT 1 FROM review_sessions WHERE id=?', (session_id,)).fetchone():
                        break
                    latest=conn.execute('SELECT payload_json,status FROM review_learning_events WHERE id=?',(event['id'],)).fetchone()
                    if not latest or latest['payload_json']!=event['payload_json'] or latest['status']!='running':
                        continue
                    current = conn.execute('SELECT version FROM review_memory WHERE session_id=?', (session_id,)).fetchone()
                    if (current['version'] if current else 0) != old['version']:
                        raise ValueError('规范在学习期间已被修改，请重试学习')
                    if pending:
                        save_questions(conn,dict(event),proposal,feedback)
                    else:
                        if has_updates:
                            save_revision(conn,session_id,old['version'],old['rules'],'AI学习前')
                            conn.execute('INSERT INTO review_memory VALUES (?, ?, ?, ?) ON CONFLICT(session_id) DO UPDATE SET '
                                         'version=excluded.version, rules_json=excluded.rules_json, updated_at=excluded.updated_at',
                                         (session_id, old['version'] + 1, dumps_json(rules), utc_now()))
                        conn.execute("UPDATE review_learning_events SET status='completed', updated_at=? WHERE id=?", (utc_now(), event['id']))
                from .session_store import add_message, emit_event
                if pending:
                    add_message(session_id,'assistant','learning_notice','这次反馈的学习范围需要确认，请在下方选择解释；不影响审校继续执行。')
                elif has_updates:
                    add_message(session_id,'assistant','learning_confirmation',f"会话规范已更新至 v{old['version']+1}。\n"+'\n'.join(rule['text'] for rule in rules if rule['id'] not in {r['id'] for r in old['rules']}))
                emit_event(session_id,'learning.updated',{'event_id':event['id'],'needs_input':pending})
                _add_log(event['task_id'], 'info', '学习等待聊天确认' if pending else '自主学习完成')
            except Exception as exc:
                with get_connection() as conn:
                    conn.execute("UPDATE review_learning_events SET status='failed', error=?, updated_at=? WHERE id=?",
                                 (str(exc), utc_now(), event['id']))
                _add_log(event['task_id'], 'error', '自主学习失败，人工反馈已保留：' + str(exc))
    finally:
        with _worker_lock:
            _workers.discard(session_id)
        # Cover a submission arriving after the final SELECT but before worker removal.
        with get_connection() as conn:
            queued = conn.execute("SELECT 1 FROM review_learning_events WHERE session_id=? AND status='queued'", (session_id,)).fetchone()
        if queued:
            start_learning(session_id)
